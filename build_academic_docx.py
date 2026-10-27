import os
import sys
import re
import shutil
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def style_table(table, header_bg="1E293B", alt_bg="F8FAFC"):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(table.rows):
        is_header = (i == 0)
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
        if is_header:
            trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
            
        for cell in row.cells:
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
            
            tcPr = cell._tc.get_or_add_tcPr()
            borders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:top w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/><w:bottom w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/><w:left w:val="none"/><w:right w:val="none"/></w:tcBorders>')
            tcPr.append(borders)
            
            if is_header:
                set_cell_background(cell, header_bg)
                for p in cell.paragraphs:
                    p.paragraph_format.space_after = Pt(2)
                    p.paragraph_format.space_before = Pt(2)
                    for r in p.runs:
                        r.font.bold = True
                        r.font.size = Pt(8.5)
                        r.font.color.rgb = RGBColor(255, 255, 255)
                        r.font.name = "Calibri"
            else:
                if i % 2 == 1:
                    set_cell_background(cell, "FFFFFF")
                else:
                    set_cell_background(cell, alt_bg)
                for p in cell.paragraphs:
                    p.paragraph_format.space_after = Pt(2)
                    p.paragraph_format.space_before = Pt(2)
                    for r in p.runs:
                        r.font.size = Pt(8.5)
                        r.font.color.rgb = RGBColor(30, 41, 59)
                        r.font.name = "Calibri"

def add_callout(doc, text, fill_hex="F8FAFC", border_color="3B82F6"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, fill_hex)
    set_cell_margins(cell, top=120, bottom=120, left=160, right=160)
    
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/><w:top w:val="none"/><w:right w:val="none"/><w:bottom w:val="none"/></w:tcBorders>')
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(text)
    r.font.name = "Consolas" if ("=" in text or "+" in text or "\\" in text) else "Calibri"
    r.font.size = Pt(9.0)
    r.font.color.rgb = RGBColor(30, 41, 59)
    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_after = Pt(2)

def append_markdown_content_to_doc(doc, content):
    lines = content.split("\n")
    i = 0
    in_code_block = False
    code_buffer = []
    in_table = False
    table_lines = []
    
    while i < len(lines):
        line = lines[i].rstrip()
        
        # Check code block
        if line.startswith("```"):
            if in_code_block:
                in_code_block = False
                code_text = "\n".join(code_buffer)
                add_callout(doc, code_text, fill_hex="0F172A", border_color="0EA5E9")
                code_buffer = []
            else:
                in_code_block = True
                code_buffer = []
            i += 1
            continue
            
        if in_code_block:
            code_buffer.append(line)
            i += 1
            continue
            
        # Check table
        if line.startswith("|") and line.endswith("|"):
            in_table = True
            table_lines.append(line)
            i += 1
            continue
        elif in_table:
            in_table = False
            raw_rows = []
            for t_line in table_lines:
                if re.match(r"^\|[\s\-:|]+\|$", t_line):
                    continue
                cells = [c.strip() for c in t_line.split("|")[1:-1]]
                if cells:
                    raw_rows.append(cells)
                    
            if raw_rows:
                num_cols = max(len(r) for r in raw_rows)
                tbl = doc.add_table(rows=len(raw_rows), cols=num_cols)
                for r_idx, row_data in enumerate(raw_rows):
                    for c_idx, cell_data in enumerate(row_data):
                        clean_cell = cell_data.replace("**", "").replace("*", "").replace("`", "")
                        tbl.cell(r_idx, c_idx).text = clean_cell
                style_table(tbl)
                doc.add_paragraph().paragraph_format.space_after = Pt(4)
                
            table_lines = []
            
        if not line:
            i += 1
            continue
            
        # Title (# )
        if line.startswith("# ") and not line.startswith("## "):
            title_text = line[2:].strip()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(8)
            r = p.add_run(title_text)
            r.font.name = "Calibri"
            r.font.size = Pt(17)
            r.font.bold = True
            r.font.color.rgb = RGBColor(15, 23, 42)
            i += 1
            continue
            
        # Section Heading (## )
        if line.startswith("## "):
            sec_text = line[3:].strip()
            h = doc.add_heading(sec_text, level=1)
            h.paragraph_format.space_before = Pt(14)
            h.paragraph_format.space_after = Pt(4)
            h.paragraph_format.keep_with_next = True
            for r in h.runs:
                r.font.name = "Calibri"
                r.font.size = Pt(13)
                r.font.bold = True
                r.font.color.rgb = RGBColor(15, 23, 42)
            i += 1
            continue
            
        # Sub-heading (### )
        if line.startswith("### "):
            sub_text = line[4:].strip()
            h = doc.add_heading(sub_text, level=2)
            h.paragraph_format.space_before = Pt(10)
            h.paragraph_format.space_after = Pt(3)
            h.paragraph_format.keep_with_next = True
            for r in h.runs:
                r.font.name = "Calibri"
                r.font.size = Pt(11)
                r.font.bold = True
                r.font.color.rgb = RGBColor(30, 58, 138)
            i += 1
            continue

        # Sub-sub-heading (#### )
        if line.startswith("#### "):
            sub_text = line[5:].strip()
            h = doc.add_heading(sub_text, level=3)
            h.paragraph_format.space_before = Pt(8)
            h.paragraph_format.space_after = Pt(2)
            h.paragraph_format.keep_with_next = True
            for r in h.runs:
                r.font.name = "Calibri"
                r.font.size = Pt(10)
                r.font.bold = True
                r.font.color.rgb = RGBColor(51, 65, 85)
            i += 1
            continue

        # Horizontal Rule
        if line.startswith("---"):
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.space_before = Pt(6)
            r = p.add_run("―" * 45)
            r.font.color.rgb = RGBColor(203, 213, 225)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            i += 1
            continue

        # Equation or math block ($$ ... $$)
        if line.startswith("$$") and line.endswith("$$") and len(line) > 2:
            eq_text = line[2:-2].strip()
            add_callout(doc, eq_text, fill_hex="F1F5F9", border_color="2563EB")
            i += 1
            continue

        # Bullet list (- or *)
        if line.startswith("- ") or line.startswith("* "):
            bullet_text = line[2:].strip()
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.15
            m_bold = re.match(r"^\*\*(.*?)\*\*(.*)$", bullet_text)
            if m_bold:
                r_b = p.add_run(m_bold.group(1))
                r_b.font.name = "Calibri"
                r_b.font.size = Pt(9.5)
                r_b.font.bold = True
                r_b.font.color.rgb = RGBColor(15, 23, 42)
                r_rem = p.add_run(m_bold.group(2).replace("**", "").replace("`", ""))
                r_rem.font.name = "Calibri"
                r_rem.font.size = Pt(9.5)
                r_rem.font.color.rgb = RGBColor(30, 41, 59)
            else:
                r = p.add_run(bullet_text.replace("**", "").replace("`", ""))
                r.font.name = "Calibri"
                r.font.size = Pt(9.5)
                r.font.color.rgb = RGBColor(30, 41, 59)
            i += 1
            continue

        # Numbered list (1. 2. etc.)
        m_num = re.match(r"^(\d+)\.\s+(.*)$", line)
        if m_num:
            num_val = m_num.group(1)
            item_text = m_num.group(2).strip()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.15
            r_num = p.add_run(f"{num_val}. ")
            r_num.font.bold = True
            r_num.font.name = "Calibri"
            r_num.font.size = Pt(9.5)
            r_num.font.color.rgb = RGBColor(30, 58, 138)
            
            m_b = re.match(r"^\*\*(.*?)\*\*(.*)$", item_text)
            if m_b:
                r_b = p.add_run(m_b.group(1))
                r_b.font.name = "Calibri"
                r_b.font.size = Pt(9.5)
                r_b.font.bold = True
                r_b.font.color.rgb = RGBColor(15, 23, 42)
                r_rem = p.add_run(m_b.group(2).replace("**", "").replace("`", ""))
                r_rem.font.name = "Calibri"
                r_rem.font.size = Pt(9.5)
                r_rem.font.color.rgb = RGBColor(30, 41, 59)
            else:
                r = p.add_run(item_text.replace("**", "").replace("`", ""))
                r.font.name = "Calibri"
                r.font.size = Pt(9.5)
                r.font.color.rgb = RGBColor(30, 41, 59)
            i += 1
            continue

        # Regular paragraph
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.line_spacing = 1.15
        
        # Check if author line or affiliation
        if "Battiprolu Karthik" in line:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(line.replace("**", ""))
            r.font.name = "Calibri"
            r.font.size = Pt(11)
            r.font.bold = True
            r.font.color.rgb = RGBColor(15, 23, 42)
            i += 1
            continue
        elif "Chaitanya Bharathi" in line or "Department of Artificial" in line or "Guide / Lab" in line:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(line.replace("**", ""))
            r.font.name = "Calibri"
            r.font.size = Pt(9.5)
            r.font.italic = True
            r.font.color.rgb = RGBColor(71, 85, 105)
            i += 1
            continue
            
        clean_p = line.replace("`", "")
        parts = re.split(r"(\*\*.*?\*\*)", clean_p)
        for part in parts:
            if part.startswith("**") and part.endswith("**") and len(part) >= 4:
                r = p.add_run(part[2:-2])
                r.font.name = "Calibri"
                r.font.size = Pt(9.5)
                r.font.bold = True
                r.font.color.rgb = RGBColor(15, 23, 42)
            else:
                r = p.add_run(part)
                r.font.name = "Calibri"
                r.font.size = Pt(9.5)
                r.font.color.rgb = RGBColor(30, 41, 59)
        i += 1

def parse_markdown_to_docx(md_path, docx_path):
    with open(md_path, "r", encoding="utf-8") as f:
        content = f.read()

    doc = docx.Document()
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)
        
    append_markdown_content_to_doc(doc, content)
    doc.save(docx_path)
    print(f"Generated: {docx_path} ({os.path.getsize(docx_path)} bytes)")

def build_consolidated_master_report(p1_md, p2_md, out_path):
    doc = docx.Document()
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Master Cover Page
    p_badge = doc.add_paragraph()
    p_badge.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_badge.paragraph_format.space_before = Pt(40)
    p_badge.paragraph_format.space_after = Pt(12)
    r_badge = p_badge.add_run("ACADEMIC CASE STUDY & RESEARCH REPORT COMPENDIUM")
    r_badge.font.name = "Calibri"
    r_badge.font.size = Pt(11)
    r_badge.font.bold = True
    r_badge.font.color.rgb = RGBColor(37, 99, 235)

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_after = Pt(12)
    r_title = p_title.add_run("VoicePrint: Deep Neural Speaker Biometric Authentication, Identification & Comparison")
    r_title.font.name = "Calibri"
    r_title.font.size = Pt(22)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(15, 23, 42)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(36)
    r_sub = p_sub.add_run("Comprehensive Research Portfolio & Laboratory Documentation\nSubject: Advanced Artificial Neural Networks (Deep Learning Laboratory)")
    r_sub.font.name = "Calibri"
    r_sub.font.size = Pt(12)
    r_sub.font.italic = True
    r_sub.font.color.rgb = RGBColor(71, 85, 105)

    p_auth = doc.add_paragraph()
    p_auth.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_auth.paragraph_format.space_after = Pt(6)
    r_auth = p_auth.add_run("Principal Investigator: BATTIPROLU KARTHIK")
    r_auth.font.name = "Calibri"
    r_auth.font.size = Pt(13)
    r_auth.font.bold = True
    r_auth.font.color.rgb = RGBColor(15, 23, 42)

    p_aff = doc.add_paragraph()
    p_aff.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_aff.paragraph_format.space_after = Pt(48)
    r_aff = p_aff.add_run("Department of Artificial Intelligence & Machine Learning\nChaitanya Bharathi Institute of Technology (CBIT), Hyderabad, India")
    r_aff.font.name = "Calibri"
    r_aff.font.size = Pt(10.5)
    r_aff.font.color.rgb = RGBColor(100, 116, 139)

    # Executive Summary Card
    add_callout(doc, 
        "EXECUTIVE SUMMARY:\nThis academic project compendium compiles two exhaustive research papers documenting VoicePrint, an enterprise-grade deep neural voice biometric system. Research Piece 1 addresses the core ECAPA-TDNN acoustic feature learning pipeline, Channel-Dependent Attentive Statistics Pooling (Ch-ASP), AAM-Softmax deep metric loss, VoxCeleb benchmark evaluation (0.83% EER), and full-stack FastAPI/React system implementation. Research Piece 2 addresses open-set 1:N speaker identification, centroid outlier rejection using Interquartile Range (IQR) disparity filtering (recovering intra-class similarity from 0.506 to 0.845), Neyman-Pearson decision boundary calibration (tau = 0.70), acoustic anti-spoofing defense, and AES-256-GCM cryptographic biometric template protection.", 
        fill_hex="F8FAFC", border_color="2563EB")

    doc.add_page_break()

    # Part I: Research Paper 1
    p_part1 = doc.add_paragraph()
    p_part1.paragraph_format.space_before = Pt(24)
    p_part1.paragraph_format.space_after = Pt(12)
    r_p1 = p_part1.add_run("PART I: FOUNDATIONAL ACOUSTIC MODEL & FULL-STACK VERIFICATION")
    r_p1.font.name = "Calibri"
    r_p1.font.size = Pt(13)
    r_p1.font.bold = True
    r_p1.font.color.rgb = RGBColor(37, 99, 235)

    with open(p1_md, "r", encoding="utf-8") as f:
        c1 = f.read()
    append_markdown_content_to_doc(doc, c1)

    doc.add_page_break()

    # Part II: Research Paper 2
    p_part2 = doc.add_paragraph()
    p_part2.paragraph_format.space_before = Pt(24)
    p_part2.paragraph_format.space_after = Pt(12)
    r_p2 = p_part2.add_run("PART II: OPEN-SET IDENTIFICATION, CENTROID OUTLIER REJECTION & BIOMETRIC DEFENSE")
    r_p2.font.name = "Calibri"
    r_p2.font.size = Pt(13)
    r_p2.font.bold = True
    r_p2.font.color.rgb = RGBColor(37, 99, 235)

    with open(p2_md, "r", encoding="utf-8") as f:
        c2 = f.read()
    append_markdown_content_to_doc(doc, c2)

    doc.save(out_path)
    print(f"Generated Consolidated Master Report: {out_path} ({os.path.getsize(out_path)} bytes)")

if __name__ == "__main__":
    p1_md = "VOICEPRINT_RESEARCH_PAPER_1.md"
    p2_md = "VOICEPRINT_RESEARCH_PAPER_2.md"
    
    p1_docx = "VoicePrint_Research_Paper_1.docx"
    p2_docx = "VoicePrint_Research_Paper_2.docx"
    master_docx = "VoicePrint_Complete_Academic_Project_Report.docx"
    
    parse_markdown_to_docx(p1_md, p1_docx)
    parse_markdown_to_docx(p2_md, p2_docx)
    build_consolidated_master_report(p1_md, p2_md, master_docx)
    
    # Copy all docx to frontend/public/
    for fn in [p1_docx, p2_docx, master_docx]:
        shutil.copyfile(fn, os.path.join("frontend", "public", fn))
        if os.path.exists("frontend/dist"):
            shutil.copyfile(fn, os.path.join("frontend", "dist", fn))
            
    print("All Word document deliverables successfully compiled and mirrored into frontend public & dist!")
